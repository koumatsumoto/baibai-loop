from __future__ import annotations

import io
import zipfile
from collections import defaultdict
from collections.abc import Sequence
from datetime import date, datetime, time

import openpyxl

from baibai_engine.collection.economic_sources import SourceDefinition

from ..observations import NativeRecords
from .base import (
    MAX_ZIP_RESPONSE_BYTES,
    FetchContext,
    HttpSession,
    IndicatorsProviderError,
    ProviderSpec,
    fetch_bytes,
)


class BojSource:
    """Bank of Japan long time-series Excel workbook (no auth).

    BOJ publishes stable ``.xlsx`` long-series workbooks (e.g. the monetary base
    at ``other/mb/mblong.xlsx``); the stat-search interactive CSV is session /
    cgi based and not a stable GET URL. The first sheet holds the headline series
    with a month-end date in the date column and values in numbered data columns.
    Full rows, leading metadata and original date cells are retained. The monthly
    date key is only a lookup coordinate. L2 validates the requested column/header/
    unit and interprets its numeric or missing value.
    """

    spec = ProviderSpec(name="boj", all_history_start=date(1957, 1, 1))
    name = spec.name

    def fetch_records(
        self,
        series: SourceDefinition,
        *,
        start: date,
        end: date,
        session: HttpSession,
        context: FetchContext | None = None,
    ) -> NativeRecords:
        content = fetch_bytes(
            session,
            series.source_url,
            params=None,
            max_bytes=MAX_ZIP_RESPONSE_BYTES,
            context=context,
        )
        return parse_boj_xlsx(series, content, start=start, end=end)


def parse_boj_xlsx(
    _series: SourceDefinition, content: bytes, *, start: date, end: date
) -> NativeRecords:
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise IndicatorsProviderError("BOJ response is not a .xlsx (zip) workbook")
    records: dict[date, list[dict[str, object]]] = defaultdict(list)
    preamble: list[list[object]] = []
    saw_date_row = False
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            worksheet = workbook[workbook.sheetnames[0]]
            for row in worksheet.iter_rows(values_only=True):
                cells: Sequence[object] = row
                observed_at = _row_month(cells)
                native_cells = [_native_cell(cell) for cell in cells]
                if observed_at is None:
                    if not saw_date_row:
                        preamble.append(native_cells)
                    continue
                saw_date_row = True
                if start <= observed_at <= end:
                    records[observed_at].append(
                        {
                            "sheet": worksheet.title,
                            "preamble": list(preamble),
                            "row": native_cells,
                        }
                    )
        finally:
            workbook.close()
    except Exception as exc:
        raise IndicatorsProviderError(f"BOJ workbook could not be read: {exc}") from exc
    if not saw_date_row:
        raise IndicatorsProviderError("BOJ workbook has no parseable date rows")
    return dict(records)


def _native_cell(value: object) -> object:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    raise ValueError(f"unsupported BOJ workbook cell type: {type(value).__name__}")


def _row_month(cells: Sequence[object]) -> date | None:
    for cell in cells:
        if isinstance(cell, datetime):
            return date(cell.year, cell.month, 1)
        if isinstance(cell, date):
            return date(cell.year, cell.month, 1)
    return None
