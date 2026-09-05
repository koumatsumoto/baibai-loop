"""Produce point-in-time TSE policy and EDINET filing context for Research Triage.

The context is a fact-layer annotation and never affects ranking, E[r], or a gate.
Filing absence is reported only when identity coverage proves the whole lookback window.
"""

from __future__ import annotations

import re
import sqlite3
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Literal

import openpyxl
import requests

from baibai_engine.market.sqlite import connect_current, open_connection
from baibai_engine.market.ticker import normalize_ticker

TSE_CAPITAL_POLICY_URL = "https://www.jpx.co.jp/equities/follow-up/jr4eth0000004vj2-att/list.xlsx"

# Six months of filings. A large-holding position or a tender offer older than that is
# no longer the "recent" context the research_triage reader is asking about, and the window
# has to be short enough that the identity coverage behind it is provable.
FILING_LOOKBACK_DAYS = 183

# EDINET form codes, from the document-type table in
# docs/reference/screening-runtime.md. Large-holding filings name their target in
# `issuerEdinetCode`; tender-offer filings name theirs in `subjectEdinetCode`.
LARGE_HOLDING_DOC_TYPES: tuple[str, ...] = ("350", "360")
TENDER_OFFER_DOC_TYPES: tuple[str, ...] = ("240", "250", "260", "270", "280")
TENDER_OFFER_REGISTRATION_DOC_TYPE = "240"
TENDER_OFFER_REGISTRATION_CORRECTION_DOC_TYPE = "250"
TENDER_OFFER_WITHDRAWAL_DOC_TYPE = "260"
TENDER_OFFER_RESULT_DOC_TYPE = "270"
TENDER_OFFER_RESULT_CORRECTION_DOC_TYPE = "280"

_HTTP_TIMEOUT_SECONDS = 30
_MONTH_SHEET_RE = re.compile(r"(?:【過去分】)?開示企業一覧（(\d{4})年(\d{1,2})月末時点）")

type CapitalPolicyStatus = Literal["disclosed", "considering"]
type ControlFilingKind = Literal["large_holding", "tender_offer"]


class ValuationCatalystError(RuntimeError):
    """A primary-source payload cannot be converted without guessing."""


@dataclass(frozen=True, slots=True, kw_only=True)
class TSECapitalPolicyRow:
    snapshot_month_end: date
    ticker: str
    status: CapitalPolicyStatus
    status_change: str | None
    updated_on: date | None
    contact_requested: bool
    first_disclosed_month_end: date | None
    first_disclosure_left_censored: bool


@dataclass(frozen=True, slots=True, kw_only=True)
class TSECapitalPolicyRefreshSummary:
    sheet_count: int
    row_count: int
    disclosed_count: int
    considering_count: int


@dataclass(frozen=True, slots=True, kw_only=True)
class ValuationCatalystContext:
    """What the sources say about one ticker, with "not observed" kept distinct.

    `None` on a field means the source could not answer for this as-of date, which is a
    different statement from "the source answered and there was nothing".
    """

    tse_capital_policy_status: Literal["disclosed", "considering", "none"] | None
    tse_capital_policy_updated_on: date | None
    large_holding_filing_within_lookback: bool | None
    latest_large_holding_filing_date: date | None
    tender_offer_filing_within_lookback: bool | None
    latest_tender_offer_filing_date: date | None


@dataclass(frozen=True, slots=True, kw_only=True)
class ControlFilingIndex:
    """The filings observed over one window, with the limits of that observation.

    Two things stop an absence from being provable. A ticker whose EDINET code the
    filing history never revealed cannot be looked up at all, so no filing could have
    been attributed to it. And a filing that named no target company could have been
    about anyone, which makes every absence of that filing kind unprovable for the whole
    window while leaving the other type answerable.
    """

    latest_by_target: Mapping[tuple[str, ControlFilingKind], date]
    identified_tickers: frozenset[str]
    anonymous_filing_kinds: frozenset[ControlFilingKind]

    def can_answer(self, ticker: str, filing_kind: ControlFilingKind) -> bool:
        return ticker in self.identified_tickers and filing_kind not in self.anonymous_filing_kinds


def parse_tse_capital_policy_workbook(content: bytes) -> tuple[TSECapitalPolicyRow, ...]:
    """Expand every monthly sheet without inventing an exact first disclosure day.

    The workbook carries the current month plus past-month snapshots, so the series is
    point-in-time reconstructable. The earliest sheet cannot say whether a company that
    already appears on it disclosed that month or earlier, so those rows are marked
    left-censored instead of being dated to the first sheet.
    """
    try:
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValuationCatalystError("failed to parse TSE capital-policy workbook") from exc

    raw_rows: list[tuple[date, str, CapitalPolicyStatus, str | None, date | None, bool]] = []
    sheet_months: list[date] = []
    for worksheet in workbook.worksheets:
        match = _MONTH_SHEET_RE.search(worksheet.title.strip())
        if match is None:
            continue
        month_end = _calendar_month_end(int(match.group(1)), int(match.group(2)))
        sheet_months.append(month_end)
        rows = list(worksheet.iter_rows(values_only=True))
        columns = _resolve_tse_columns(rows, sheet=worksheet.title)
        for values in rows[columns.first_data_row :]:
            ticker_raw = _cell(values, columns.ticker)
            status_raw = _text(_cell(values, columns.status))
            if ticker_raw in (None, "") and not status_raw:
                continue
            raw_rows.append(
                (
                    month_end,
                    _ticker_from_excel(ticker_raw),
                    _capital_policy_status(status_raw),
                    _text(_cell(values, columns.status_change)) or None,
                    _excel_date(_cell(values, columns.updated_on)),
                    bool(_text(_cell(values, columns.contact_requested))),
                )
            )
    if not sheet_months or not raw_rows:
        raise ValuationCatalystError("TSE capital-policy workbook has no monthly rows")
    if len(sheet_months) != len(set(sheet_months)):
        raise ValuationCatalystError("TSE capital-policy workbook repeats a snapshot month")

    earliest_sheet = min(sheet_months)
    first_disclosed: dict[str, date] = {}
    for month_end, ticker, status, *_ in sorted(raw_rows):
        if status == "disclosed" and ticker not in first_disclosed:
            first_disclosed[ticker] = month_end

    return tuple(
        TSECapitalPolicyRow(
            snapshot_month_end=month_end,
            ticker=ticker,
            status=status,
            status_change=status_change,
            updated_on=updated_on,
            contact_requested=contact_requested,
            first_disclosed_month_end=first_disclosed.get(ticker),
            first_disclosure_left_censored=first_disclosed.get(ticker) == earliest_sheet,
        )
        for month_end, ticker, status, status_change, updated_on, contact_requested in sorted(
            raw_rows
        )
    )


def store_tse_capital_policy_rows(sqlite_path: Path, rows: Sequence[TSECapitalPolicyRow]) -> int:
    """Replace whole months, because a month's sheet is a complete snapshot of it."""
    if not rows:
        raise ValuationCatalystError("refusing to replace TSE capital-policy facts with zero rows")
    month_ends = sorted({row.snapshot_month_end for row in rows})
    connection = open_connection(sqlite_path)
    try:
        placeholders = ",".join("?" for _ in month_ends)
        connection.execute(
            "DELETE FROM tse_capital_policy_snapshots "
            f"WHERE snapshot_month_end IN ({placeholders})",  # nosec B608
            tuple(day.isoformat() for day in month_ends),
        )
        connection.executemany(
            "INSERT INTO tse_capital_policy_snapshots("
            "snapshot_month_end, ticker, status, status_change, updated_on, "
            "contact_requested, first_disclosed_month_end, first_disclosure_left_censored"
            ") VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            (
                (
                    row.snapshot_month_end.isoformat(),
                    row.ticker,
                    row.status,
                    row.status_change,
                    _date_iso(row.updated_on),
                    int(row.contact_requested),
                    _date_iso(row.first_disclosed_month_end),
                    int(row.first_disclosure_left_censored),
                )
                for row in rows
            ),
        )
        connection.commit()
        return len(rows)
    finally:
        connection.close()


def refresh_tse_capital_policy(
    sqlite_path: Path, *, session: requests.Session | None = None
) -> TSECapitalPolicyRefreshSummary:
    """Download and replace every month present in the TSE workbook."""
    client = session or requests.Session()
    response = client.get(TSE_CAPITAL_POLICY_URL, timeout=_HTTP_TIMEOUT_SECONDS)
    if response.status_code >= 400:
        raise ValuationCatalystError(
            f"failed to download TSE capital-policy workbook: {response.status_code}"
        )
    rows = parse_tse_capital_policy_workbook(response.content)
    store_tse_capital_policy_rows(sqlite_path, rows)
    return TSECapitalPolicyRefreshSummary(
        sheet_count=len({row.snapshot_month_end for row in rows}),
        row_count=len(rows),
        disclosed_count=sum(row.status == "disclosed" for row in rows),
        considering_count=sum(row.status == "considering" for row in rows),
    )


def read_valuation_catalyst_contexts(
    sqlite_path: Path, *, as_of: date, tickers: Iterable[str]
) -> dict[str, ValuationCatalystContext]:
    """Read point-in-time policy state and recent target-company filings.

    Never affects rank: the caller copies these into the Security Analysis payload only.
    """
    requested = sorted({normalize_ticker(ticker) for ticker in tickers})
    if not requested:
        return {}
    connection = connect_current(sqlite_path)
    if connection is None:
        return {}
    try:
        policy_month = connection.execute(
            "SELECT MAX(snapshot_month_end) FROM tse_capital_policy_snapshots "
            "WHERE snapshot_month_end <= ?",
            (as_of.isoformat(),),
        ).fetchone()[0]
        filing_floor = as_of - timedelta(days=FILING_LOOKBACK_DAYS)
        filing_index = read_control_filing_index(connection, start=filing_floor, end=as_of)
        policy_by_ticker = _policy_by_ticker(connection, policy_month, requested)
    finally:
        connection.close()
    result: dict[str, ValuationCatalystContext] = {}
    for ticker in requested:
        policy = policy_by_ticker.get(ticker)
        large_recent, large_on = _filing_answer(filing_index, ticker, "large_holding")
        tender_recent, tender_on = _filing_answer(filing_index, ticker, "tender_offer")
        result[ticker] = ValuationCatalystContext(
            tse_capital_policy_status=(
                None if policy_month is None else (policy[0] if policy is not None else "none")
            ),
            tse_capital_policy_updated_on=None if policy is None else policy[1],
            large_holding_filing_within_lookback=large_recent,
            latest_large_holding_filing_date=large_on,
            tender_offer_filing_within_lookback=tender_recent,
            latest_tender_offer_filing_date=tender_on,
        )
    return result


def _filing_answer(
    index: ControlFilingIndex | None, ticker: str, filing_kind: ControlFilingKind
) -> tuple[bool | None, date | None]:
    """Answer only when absence is provable for this ticker and filing kind."""
    if index is None or not index.can_answer(ticker, filing_kind):
        return None, None
    latest = index.latest_by_target.get((ticker, filing_kind))
    return latest is not None, latest


def read_control_filing_index(
    connection: sqlite3.Connection, *, start: date, end: date
) -> ControlFilingIndex | None:
    """Latest large-holding / tender-offer filing per target ticker in the window.

    Returns ``None`` when the window is not fully identity-covered, so that the caller
    reports "not observed" rather than "no filings".
    """
    if not edinet_identity_covered(connection, start=start, end=end):
        return None
    ticker_by_edinet_code = edinet_code_to_ticker(connection)
    rows = connection.execute(
        "SELECT doc_date, doc_type_code, issuer_edinet_code, subject_edinet_code, "
        "legal_status, disclosure_status, withdrawal_status "
        "FROM edinet_documents WHERE doc_date BETWEEN ? AND ? "
        "AND doc_type_code IN "
        f"({_placeholders(LARGE_HOLDING_DOC_TYPES + TENDER_OFFER_DOC_TYPES)})",  # nosec B608
        (start.isoformat(), end.isoformat(), *LARGE_HOLDING_DOC_TYPES, *TENDER_OFFER_DOC_TYPES),
    ).fetchall()
    index: dict[tuple[str, ControlFilingKind], date] = {}
    anonymous: set[ControlFilingKind] = set()
    for doc_date, doc_type, issuer, subject, legal, disclosure, withdrawal in rows:
        if not is_usable_filing_status(legal, disclosure, withdrawal):
            continue
        filing_kind: ControlFilingKind = (
            "large_holding" if str(doc_type) in LARGE_HOLDING_DOC_TYPES else "tender_offer"
        )
        target = issuer if filing_kind == "large_holding" else subject
        if target in (None, ""):
            # The filing named no target company at all, so it could have been about any
            # ticker. That makes an absence unprovable for this filing kind over this
            # window, while the other type stays answerable.
            anonymous.add(filing_kind)
            continue
        ticker = ticker_by_edinet_code.get(str(target))
        if ticker is None:
            # The target is named but its listing cannot be resolved. That is a gap for
            # the company behind that code, which `identified_tickers` already excludes.
            continue
        observed = date.fromisoformat(str(doc_date))
        key = (ticker, filing_kind)
        if observed > index.get(key, date.min):
            index[key] = observed
    return ControlFilingIndex(
        latest_by_target=index,
        identified_tickers=frozenset(ticker_by_edinet_code.values()),
        anonymous_filing_kinds=frozenset(anonymous),
    )


def edinet_identity_covered(connection: sqlite3.Connection, *, start: date, end: date) -> bool:
    """Whether every day in the window was listed and carries submitter identity.

    Two things can leave a hole. A day the store never listed has no rows at all, and
    counting filings over it would silently read zero. A day listed before the identity
    columns existed has typed, currently inspectable rows whose `edinetCode` is null;
    the API supplies that field for every such document, so a null one is a stale row
    rather than a filing without a submitter. Expired rows cannot regain identity on a
    later list and are unusable as events. The list also carries untyped operation rows
    that legitimately have neither a document type nor a submitter, and those are not
    evidence either way.
    """
    listed_days = int(
        connection.execute(
            "SELECT COUNT(*) FROM edinet_document_lists WHERE doc_date BETWEEN ? AND ?",
            (start.isoformat(), end.isoformat()),
        ).fetchone()[0]
        or 0
    )
    if listed_days != (end - start).days + 1:
        return False
    stale_rows = int(
        connection.execute(
            "SELECT COUNT(*) FROM edinet_documents WHERE doc_date BETWEEN ? AND ? "
            "AND doc_type_code IS NOT NULL AND edinet_code IS NULL "
            "AND legal_status IN ('1', '2')",
            (start.isoformat(), end.isoformat()),
        ).fetchone()[0]
        or 0
    )
    return stale_rows == 0


def edinet_code_to_ticker(connection: sqlite3.Connection) -> dict[str, str]:
    """Resolve EDINET codes through the (edinetCode, secCode) pairs the list observed.

    A code seen against more than one ticker is dropped rather than picked between.
    """
    observed: dict[str, set[str]] = {}
    for edinet_code, sec_code in connection.execute(
        "SELECT DISTINCT edinet_code, sec_code FROM edinet_documents "
        "WHERE edinet_code IS NOT NULL AND sec_code IS NOT NULL"
    ):
        try:
            ticker = normalize_ticker(str(sec_code)[:4])
        except ValueError:
            continue
        observed.setdefault(str(edinet_code), set()).add(ticker)
    return {code: next(iter(values)) for code, values in observed.items() if len(values) == 1}


def is_usable_filing_status(legal: object, disclosure: object, withdrawal: object) -> bool:
    """Same usability rule the metric extraction applies to a list row."""
    return (
        _status_text(legal) in (None, "1", "2")
        and _status_text(disclosure) in (None, "0")
        and _status_text(withdrawal) in (None, "0")
    )


def edinet_filing_counts(sqlite_path: Path, *, through: date) -> dict[str, int]:
    """Form-code histogram of the stored index, for the source-coverage record."""
    connection = connect_current(sqlite_path)
    if connection is None:
        return {}
    doc_types = LARGE_HOLDING_DOC_TYPES + TENDER_OFFER_DOC_TYPES
    try:
        rows = connection.execute(
            "SELECT doc_type_code, COUNT(*) FROM edinet_documents "
            f"WHERE doc_date <= ? AND doc_type_code IN ({_placeholders(doc_types)}) "  # nosec B608
            "GROUP BY doc_type_code ORDER BY doc_type_code",
            (through.isoformat(), *doc_types),
        ).fetchall()
    finally:
        connection.close()
    return {str(code): int(count) for code, count in rows}


def _policy_by_ticker(
    connection: sqlite3.Connection, policy_month: object, requested: Sequence[str]
) -> dict[str, tuple[CapitalPolicyStatus, date | None]]:
    if policy_month is None:
        return {}
    rows = connection.execute(
        "SELECT ticker, status, updated_on FROM tse_capital_policy_snapshots "
        f"WHERE snapshot_month_end = ? AND ticker IN ({_placeholders(requested)})",  # nosec B608
        (str(policy_month), *requested),
    ).fetchall()
    return {
        str(ticker): (_capital_policy_status_from_store(status), _optional_iso_date(updated_on))
        for ticker, status, updated_on in rows
    }


@dataclass(frozen=True, slots=True, kw_only=True)
class _TSEColumns:
    """Where each field sits on one monthly sheet, resolved from its own header."""

    first_data_row: int
    ticker: int
    status: int
    status_change: int
    updated_on: int
    contact_requested: int


# The header spans two rows: a group label with merged cells above, and the label of
# each column inside the group below. Both are needed because a group label alone is
# ambiguous — the contact block prints three columns under one heading, and the column
# holding the disclosure text was inserted before that block in later sheets, so a
# fixed index reads free-form disclosure text as a contact request.
_TSE_TICKER_HEADER = "証券コード"
_TSE_STATUS_GROUP = "開示状況"
_TSE_STATUS_HEADER = "要請に基づく開示状況"
_TSE_STATUS_CHANGE_HEADER = "前月からの開示状況の変更"
_TSE_UPDATED_ON_HEADER = "開示内容のアップデート日"
_TSE_CONTACT_GROUP_PREFIX = "機関投資家からのより活発なコンタクトを希望"
_TSE_CONTACT_HEADER = "申請状況"


def _resolve_tse_columns(rows: Sequence[Sequence[object]], *, sheet: str) -> _TSEColumns:
    header_index = next(
        (
            index
            for index, values in enumerate(rows[:20])
            for labels in ({_text(value) for value in values if value is not None},)
            if _TSE_TICKER_HEADER in labels and _TSE_STATUS_GROUP in labels
        ),
        None,
    )
    if header_index is None or header_index + 1 >= len(rows):
        raise ValuationCatalystError(f"TSE capital-policy sheet has no typed header: {sheet}")
    group = _forward_filled_labels(rows[header_index])
    detail = [_text(value) for value in rows[header_index + 1]]
    width = max(len(group), len(detail))

    def find(predicate: Callable[[int], bool], label: str) -> int:
        column = next((index for index in range(width) if predicate(index)), None)
        if column is None:
            raise ValuationCatalystError(f"TSE capital-policy sheet has no {label} column: {sheet}")
        return column

    def at(labels: Sequence[str], index: int) -> str:
        return labels[index] if index < len(labels) else ""

    return _TSEColumns(
        first_data_row=header_index + 2,
        ticker=find(lambda index: at(group, index) == _TSE_TICKER_HEADER, _TSE_TICKER_HEADER),
        status=find(lambda index: at(detail, index) == _TSE_STATUS_HEADER, _TSE_STATUS_HEADER),
        status_change=find(
            lambda index: at(detail, index) == _TSE_STATUS_CHANGE_HEADER,
            _TSE_STATUS_CHANGE_HEADER,
        ),
        updated_on=find(
            lambda index: at(group, index) == _TSE_UPDATED_ON_HEADER, _TSE_UPDATED_ON_HEADER
        ),
        contact_requested=find(
            lambda index: (
                at(group, index).startswith(_TSE_CONTACT_GROUP_PREFIX)
                and at(detail, index) == _TSE_CONTACT_HEADER
            ),
            _TSE_CONTACT_GROUP_PREFIX,
        ),
    )


def _forward_filled_labels(values: Sequence[object]) -> list[str]:
    """Spread a merged group label across the columns it covers."""
    labels: list[str] = []
    current = ""
    for value in values:
        text = _text(value)
        if text:
            current = text
        labels.append(current)
    return labels


def _capital_policy_status(value: str) -> CapitalPolicyStatus:
    if value == "開示済":
        return "disclosed"
    if value == "検討中":
        return "considering"
    raise ValuationCatalystError(f"unknown TSE capital-policy status: {value!r}")


def _capital_policy_status_from_store(value: object) -> CapitalPolicyStatus:
    raw = str(value)
    if raw == "disclosed":
        return "disclosed"
    if raw == "considering":
        return "considering"
    raise ValuationCatalystError(f"invalid stored TSE capital-policy status: {raw!r}")


def _ticker_from_excel(value: object) -> str:
    if isinstance(value, bool):
        raise ValuationCatalystError(f"invalid TSE ticker: {value!r}")
    if isinstance(value, int):
        raw = str(value)
    elif isinstance(value, float) and value.is_integer():
        raw = str(int(value))
    else:
        raw = _text(value)
    try:
        return normalize_ticker(raw)
    except ValueError as exc:
        raise ValuationCatalystError(f"invalid TSE ticker: {value!r}") from exc


def _excel_date(value: object) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = _text(value)
    try:
        return date.fromisoformat(
            raw.replace("/", "-").replace("年", "-").replace("月", "-").removesuffix("日")
        )
    except ValueError as exc:
        raise ValuationCatalystError(f"invalid TSE update date: {raw!r}") from exc


def _calendar_month_end(year: int, month: int) -> date:
    first_next = date(year + int(month == 12), month % 12 + 1, 1)
    return first_next - timedelta(days=1)


def _cell(values: Sequence[object], index: int) -> object | None:
    return values[index] if index < len(values) else None


def _text(value: object) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip()


def _status_text(value: object) -> str | None:
    return None if value is None else str(value)


def _placeholders(values: Sequence[object]) -> str:
    return ",".join("?" for _ in values)


def _date_iso(value: date | None) -> str | None:
    return None if value is None else value.isoformat()


def _optional_iso_date(value: object) -> date | None:
    return None if value in (None, "") else date.fromisoformat(str(value))


__all__ = (
    "FILING_LOOKBACK_DAYS",
    "LARGE_HOLDING_DOC_TYPES",
    "TENDER_OFFER_DOC_TYPES",
    "TENDER_OFFER_REGISTRATION_CORRECTION_DOC_TYPE",
    "TENDER_OFFER_REGISTRATION_DOC_TYPE",
    "TENDER_OFFER_RESULT_CORRECTION_DOC_TYPE",
    "TENDER_OFFER_RESULT_DOC_TYPE",
    "TENDER_OFFER_WITHDRAWAL_DOC_TYPE",
    "CapitalPolicyStatus",
    "ControlFilingIndex",
    "ControlFilingKind",
    "TSECapitalPolicyRefreshSummary",
    "TSECapitalPolicyRow",
    "ValuationCatalystContext",
    "ValuationCatalystError",
    "edinet_code_to_ticker",
    "edinet_filing_counts",
    "edinet_identity_covered",
    "is_usable_filing_status",
    "parse_tse_capital_policy_workbook",
    "read_control_filing_index",
    "read_valuation_catalyst_contexts",
    "refresh_tse_capital_policy",
    "store_tse_capital_policy_rows",
)
