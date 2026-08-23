"""Typed capital-policy and control-event facts used only as annotations.

Three primary sources describe who might close a valuation gap and when: the TSE
"management conscious of cost of capital" disclosure list, the EDINET filing index for
large-holding and tender-offer events, and the JPX delisting record. None of them feeds
ranking, E[r] or any gate — event deltas measured negative over 3y/5y, and the disclosure
rate alone no longer separates companies. They exist so a human reading a shortlist can
see the dated catalyst context, and so a completed tender offer can price a delisted name
in the calibration forward window.

Absence of an event is only reported when the filing index was actually observed over the
whole window. The identity columns the index reads were added to the store after those
rows were first written, so a day that predates them reads as unobserved rather than as a
day on which nothing was filed.
"""

from __future__ import annotations

import re
import sqlite3
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from typing import Literal
from urllib.parse import urljoin

import openpyxl
import requests

from baibai_engine.market.sqlite import connect_current, open_connection
from baibai_engine.market.ticker import normalize_ticker

TSE_CAPITAL_POLICY_URL = "https://www.jpx.co.jp/equities/follow-up/jr4eth0000004vj2-att/list.xlsx"
JPX_DELISTING_INDEX_URL = "https://www.jpx.co.jp/listing/stocks/delisted/"

# Six months of filings. A large-holding position or a tender offer older than that is
# no longer the "recent" context the shortlist reader is asking about, and the window
# has to be short enough that the identity coverage behind it is provable.
EVENT_RECENT_DAYS = 183

# EDINET form codes, from the document-type table in
# docs/reference/screening-runtime.md. Large-holding filings name their target in
# `issuerEdinetCode`; tender-offer filings name theirs in `subjectEdinetCode`.
LARGE_HOLDING_DOC_TYPES: tuple[str, ...] = ("350", "360")
TENDER_OFFER_DOC_TYPES: tuple[str, ...] = ("240", "250", "260", "270", "280")
TENDER_OFFER_REGISTRATION_DOC_TYPE = "240"
TENDER_OFFER_REGISTRATION_CORRECTION_DOC_TYPE = "250"
TENDER_OFFER_WITHDRAWAL_DOC_TYPE = "260"
TENDER_OFFER_RESULT_DOC_TYPE = "270"
TENDER_OFFER_RESULT_CORRECTION_DOC_TYPE = "280"

_HTTP_TIMEOUT_SECONDS = 30
_MONTH_SHEET_RE = re.compile(r"(?:【過去分】)?開示企業一覧（(\d{4})年(\d{1,2})月末時点）")
_JPX_ARCHIVE_RE = re.compile(r"/listing/stocks/delisted/(?:index\.html|archives-\d+\.html)$")

type CapitalPolicyStatus = Literal["disclosed", "considering"]
type ControlEventType = Literal["large_holding", "tender_offer"]


class CapitalControlError(RuntimeError):
    """A primary-source payload cannot be converted without guessing."""


@dataclass(frozen=True, slots=True, kw_only=True)
class TSECapitalPolicyRow:
    snapshot_month_end: date
    ticker: str
    status: CapitalPolicyStatus
    status_change: str | None
    updated_on: date | None
    contact_requested: bool
    first_disclosed_month_end: date | None
    first_disclosure_left_censored: bool


@dataclass(frozen=True, slots=True, kw_only=True)
class JPXDelistingRow:
    delisted_on: date
    ticker: str
    name: str
    market: str | None
    reason: str


@dataclass(frozen=True, slots=True, kw_only=True)
class CapitalControlAnnotation:
    """What the sources say about one ticker, with "not observed" kept distinct.

    `None` on a field means the source could not answer for this as-of date, which is a
    different statement from "the source answered and there was nothing".
    """

    tse_capital_policy_status: Literal["disclosed", "considering", "none"] | None
    tse_capital_policy_updated_on: date | None
    large_holding_event_recent: bool | None
    large_holding_event_latest_on: date | None
    tender_offer_event_recent: bool | None
    tender_offer_event_latest_on: date | None


@dataclass(frozen=True, slots=True, kw_only=True)
class ControlEventIndex:
    """The filings observed over one window, with the limits of that observation.

    Two things stop an absence from being provable. A ticker whose EDINET code the
    filing history never revealed cannot be looked up at all, so no filing could have
    been attributed to it. And a filing that named no target company could have been
    about anyone, which makes every absence of that event type unprovable for the whole
    window while leaving the other type answerable.
    """

    latest_by_target: Mapping[tuple[str, ControlEventType], date]
    identified_tickers: frozenset[str]
    anonymous_event_types: frozenset[ControlEventType]

    def can_answer(self, ticker: str, event_type: ControlEventType) -> bool:
        return ticker in self.identified_tickers and event_type not in self.anonymous_event_types


@dataclass(frozen=True, slots=True, kw_only=True)
class CapitalControlRefreshSummary:
    tse_sheet_count: int
    tse_row_count: int
    tse_disclosed_count: int
    tse_considering_count: int
    jpx_delisting_count: int


def parse_tse_capital_policy_workbook(content: bytes) -> tuple[TSECapitalPolicyRow, ...]:
    """Expand every monthly sheet without inventing an exact first disclosure day.

    The workbook carries the current month plus past-month snapshots, so the series is
    point-in-time reconstructable. The earliest sheet cannot say whether a company that
    already appears on it disclosed that month or earlier, so those rows are marked
    left-censored instead of being dated to the first sheet.
    """
    try:
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise CapitalControlError("failed to parse TSE capital-policy workbook") from exc

    raw_rows: list[tuple[date, str, CapitalPolicyStatus, str | None, date | None, bool]] = []
    sheet_months: list[date] = []
    for worksheet in workbook.worksheets:
        match = _MONTH_SHEET_RE.search(worksheet.title.strip())
        if match is None:
            continue
        month_end = _calendar_month_end(int(match.group(1)), int(match.group(2)))
        sheet_months.append(month_end)
        rows = list(worksheet.iter_rows(values_only=True))
        columns = _resolve_tse_columns(rows, sheet=worksheet.title)
        for values in rows[columns.first_data_row :]:
            ticker_raw = _cell(values, columns.ticker)
            status_raw = _text(_cell(values, columns.status))
            if ticker_raw in (None, "") and not status_raw:
                continue
            raw_rows.append(
                (
                    month_end,
                    _ticker_from_excel(ticker_raw),
                    _capital_policy_status(status_raw),
                    _text(_cell(values, columns.status_change)) or None,
                    _excel_date(_cell(values, columns.updated_on)),
                    bool(_text(_cell(values, columns.contact_requested))),
                )
            )
    if not sheet_months or not raw_rows:
        raise CapitalControlError("TSE capital-policy workbook has no monthly rows")
    if len(sheet_months) != len(set(sheet_months)):
        raise CapitalControlError("TSE capital-policy workbook repeats a snapshot month")

    earliest_sheet = min(sheet_months)
    first_disclosed: dict[str, date] = {}
    for month_end, ticker, status, *_ in sorted(raw_rows):
        if status == "disclosed" and ticker not in first_disclosed:
            first_disclosed[ticker] = month_end

    return tuple(
        TSECapitalPolicyRow(
            snapshot_month_end=month_end,
            ticker=ticker,
            status=status,
            status_change=status_change,
            updated_on=updated_on,
            contact_requested=contact_requested,
            first_disclosed_month_end=first_disclosed.get(ticker),
            first_disclosure_left_censored=first_disclosed.get(ticker) == earliest_sheet,
        )
        for month_end, ticker, status, status_change, updated_on, contact_requested in sorted(
            raw_rows
        )
    )


def store_tse_capital_policy_rows(sqlite_path: Path, rows: Sequence[TSECapitalPolicyRow]) -> int:
    """Replace whole months, because a month's sheet is a complete snapshot of it."""
    if not rows:
        raise CapitalControlError("refusing to replace TSE capital-policy facts with zero rows")
    month_ends = sorted({row.snapshot_month_end for row in rows})
    connection = open_connection(sqlite_path)
    try:
        placeholders = ",".join("?" for _ in month_ends)
        connection.execute(
            "DELETE FROM tse_capital_policy_snapshots "
            f"WHERE snapshot_month_end IN ({placeholders})",  # nosec B608
            tuple(day.isoformat() for day in month_ends),
        )
        connection.executemany(
            "INSERT INTO tse_capital_policy_snapshots("
            "snapshot_month_end, ticker, status, status_change, updated_on, "
            "contact_requested, first_disclosed_month_end, first_disclosure_left_censored"
            ") VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            (
                (
                    row.snapshot_month_end.isoformat(),
                    row.ticker,
                    row.status,
                    row.status_change,
                    _date_iso(row.updated_on),
                    int(row.contact_requested),
                    _date_iso(row.first_disclosed_month_end),
                    int(row.first_disclosure_left_censored),
                )
                for row in rows
            ),
        )
        connection.commit()
        return len(rows)
    finally:
        connection.close()


def store_jpx_delistings(sqlite_path: Path, rows: Sequence[JPXDelistingRow]) -> int:
    """Accumulate delisting facts, because JPX retires archive pages over time.

    A delisting is immutable once it happened, so a page that no longer lists 2024 must
    not erase 2024. Re-reading a listed year overwrites its rows with the newer reading.
    """
    if not rows:
        raise CapitalControlError("refusing to store zero JPX delistings")
    connection = open_connection(sqlite_path)
    try:
        connection.executemany(
            "INSERT OR REPLACE INTO jpx_delistings(delisted_on, ticker, name, market, reason) "
            "VALUES (?, ?, ?, ?, ?)",
            (
                (row.delisted_on.isoformat(), row.ticker, row.name, row.market, row.reason)
                for row in rows
            ),
        )
        connection.commit()
        return len(rows)
    finally:
        connection.close()


def refresh_capital_control_facts(
    sqlite_path: Path, *, session: requests.Session | None = None
) -> CapitalControlRefreshSummary:
    """Re-read both JPX sources and replace what they cover. Safe to repeat."""
    client = session or requests.Session()
    workbook_response = client.get(TSE_CAPITAL_POLICY_URL, timeout=_HTTP_TIMEOUT_SECONDS)
    if workbook_response.status_code >= 400:
        raise CapitalControlError(
            f"failed to download TSE capital-policy workbook: {workbook_response.status_code}"
        )
    tse_rows = parse_tse_capital_policy_workbook(workbook_response.content)
    delistings = _download_jpx_delistings(client)
    store_tse_capital_policy_rows(sqlite_path, tse_rows)
    store_jpx_delistings(sqlite_path, delistings)
    return CapitalControlRefreshSummary(
        tse_sheet_count=len({row.snapshot_month_end for row in tse_rows}),
        tse_row_count=len(tse_rows),
        tse_disclosed_count=sum(row.status == "disclosed" for row in tse_rows),
        tse_considering_count=sum(row.status == "considering" for row in tse_rows),
        jpx_delisting_count=len(delistings),
    )


def read_jpx_delistings(sqlite_path: Path) -> tuple[JPXDelistingRow, ...]:
    connection = connect_current(sqlite_path)
    if connection is None:
        return ()
    try:
        rows = connection.execute(
            "SELECT delisted_on, ticker, name, market, reason FROM jpx_delistings "
            "ORDER BY delisted_on, ticker"
        ).fetchall()
    finally:
        connection.close()
    return tuple(
        JPXDelistingRow(
            delisted_on=date.fromisoformat(str(delisted_on)),
            ticker=str(ticker),
            name=str(name),
            market=None if market is None else str(market),
            reason=str(reason),
        )
        for delisted_on, ticker, name, market, reason in rows
    )


def read_capital_control_annotations(
    sqlite_path: Path, *, asof: date, tickers: Iterable[str]
) -> dict[str, CapitalControlAnnotation]:
    """Read point-in-time policy state and recent target-company filings.

    Never affects rank: the caller copies these into the candidate payload only.
    """
    requested = sorted({normalize_ticker(ticker) for ticker in tickers})
    if not requested:
        return {}
    connection = connect_current(sqlite_path)
    if connection is None:
        return {}
    try:
        policy_month = connection.execute(
            "SELECT MAX(snapshot_month_end) FROM tse_capital_policy_snapshots "
            "WHERE snapshot_month_end <= ?",
            (asof.isoformat(),),
        ).fetchone()[0]
        event_floor = asof - timedelta(days=EVENT_RECENT_DAYS)
        event_index = read_control_event_index(connection, start=event_floor, end=asof)
        policy_by_ticker = _policy_by_ticker(connection, policy_month, requested)
    finally:
        connection.close()
    result: dict[str, CapitalControlAnnotation] = {}
    for ticker in requested:
        policy = policy_by_ticker.get(ticker)
        large_recent, large_on = _event_answer(event_index, ticker, "large_holding")
        tender_recent, tender_on = _event_answer(event_index, ticker, "tender_offer")
        result[ticker] = CapitalControlAnnotation(
            tse_capital_policy_status=(
                None if policy_month is None else (policy[0] if policy is not None else "none")
            ),
            tse_capital_policy_updated_on=None if policy is None else policy[1],
            large_holding_event_recent=large_recent,
            large_holding_event_latest_on=large_on,
            tender_offer_event_recent=tender_recent,
            tender_offer_event_latest_on=tender_on,
        )
    return result


def _event_answer(
    index: ControlEventIndex | None, ticker: str, event_type: ControlEventType
) -> tuple[bool | None, date | None]:
    """Answer only when an absence would be provable for this ticker and event type."""
    if index is None or not index.can_answer(ticker, event_type):
        return None, None
    latest = index.latest_by_target.get((ticker, event_type))
    return latest is not None, latest


def read_control_event_index(
    connection: sqlite3.Connection, *, start: date, end: date
) -> ControlEventIndex | None:
    """Latest large-holding / tender-offer filing per target ticker in the window.

    Returns ``None`` when the window is not fully identity-covered, so that the caller
    reports "not observed" rather than "no filings".
    """
    if not edinet_identity_covered(connection, start=start, end=end):
        return None
    ticker_by_edinet_code = edinet_code_to_ticker(connection)
    rows = connection.execute(
        "SELECT doc_date, doc_type_code, issuer_edinet_code, subject_edinet_code, "
        "legal_status, disclosure_status, withdrawal_status "
        "FROM edinet_documents WHERE doc_date BETWEEN ? AND ? "
        "AND doc_type_code IN "
        f"({_placeholders(LARGE_HOLDING_DOC_TYPES + TENDER_OFFER_DOC_TYPES)})",  # nosec B608
        (start.isoformat(), end.isoformat(), *LARGE_HOLDING_DOC_TYPES, *TENDER_OFFER_DOC_TYPES),
    ).fetchall()
    index: dict[tuple[str, ControlEventType], date] = {}
    anonymous: set[ControlEventType] = set()
    for doc_date, doc_type, issuer, subject, legal, disclosure, withdrawal in rows:
        if not is_usable_filing_status(legal, disclosure, withdrawal):
            continue
        event_type: ControlEventType = (
            "large_holding" if str(doc_type) in LARGE_HOLDING_DOC_TYPES else "tender_offer"
        )
        target = issuer if event_type == "large_holding" else subject
        if target in (None, ""):
            # The filing named no target company at all, so it could have been about any
            # ticker. That makes an absence unprovable for this event type over this
            # window, while the other type stays answerable.
            anonymous.add(event_type)
            continue
        ticker = ticker_by_edinet_code.get(str(target))
        if ticker is None:
            # The target is named but its listing cannot be resolved. That is a gap for
            # the company behind that code, which `identified_tickers` already excludes.
            continue
        observed = date.fromisoformat(str(doc_date))
        key = (ticker, event_type)
        if observed > index.get(key, date.min):
            index[key] = observed
    return ControlEventIndex(
        latest_by_target=index,
        identified_tickers=frozenset(ticker_by_edinet_code.values()),
        anonymous_event_types=frozenset(anonymous),
    )


def edinet_identity_covered(connection: sqlite3.Connection, *, start: date, end: date) -> bool:
    """Whether every day in the window was listed and carries submitter identity.

    Two things can leave a hole. A day the store never listed has no rows at all, and
    counting filings over it would silently read zero. A day listed before the identity
    columns existed has typed, currently inspectable rows whose `edinetCode` is null;
    the API supplies that field for every such document, so a null one is a stale row
    rather than a filing without a submitter. Expired rows cannot regain identity on a
    later list and are unusable as events. The list also carries untyped operation rows
    that legitimately have neither a document type nor a submitter, and those are not
    evidence either way.
    """
    listed_days = int(
        connection.execute(
            "SELECT COUNT(*) FROM edinet_document_lists WHERE doc_date BETWEEN ? AND ?",
            (start.isoformat(), end.isoformat()),
        ).fetchone()[0]
        or 0
    )
    if listed_days != (end - start).days + 1:
        return False
    stale_rows = int(
        connection.execute(
            "SELECT COUNT(*) FROM edinet_documents WHERE doc_date BETWEEN ? AND ? "
            "AND doc_type_code IS NOT NULL AND edinet_code IS NULL "
            "AND legal_status IN ('1', '2')",
            (start.isoformat(), end.isoformat()),
        ).fetchone()[0]
        or 0
    )
    return stale_rows == 0


def edinet_code_to_ticker(connection: sqlite3.Connection) -> dict[str, str]:
    """Resolve EDINET codes through the (edinetCode, secCode) pairs the list observed.

    A code seen against more than one ticker is dropped rather than picked between.
    """
    observed: dict[str, set[str]] = {}
    for edinet_code, sec_code in connection.execute(
        "SELECT DISTINCT edinet_code, sec_code FROM edinet_documents "
        "WHERE edinet_code IS NOT NULL AND sec_code IS NOT NULL"
    ):
        try:
            ticker = normalize_ticker(str(sec_code)[:4])
        except ValueError:
            continue
        observed.setdefault(str(edinet_code), set()).add(ticker)
    return {code: next(iter(values)) for code, values in observed.items() if len(values) == 1}


def is_usable_filing_status(legal: object, disclosure: object, withdrawal: object) -> bool:
    """Same usability rule the metric extraction applies to a list row."""
    return (
        _status_text(legal) in (None, "1", "2")
        and _status_text(disclosure) in (None, "0")
        and _status_text(withdrawal) in (None, "0")
    )


def edinet_event_counts(sqlite_path: Path, *, through: date) -> dict[str, int]:
    """Form-code histogram of the stored index, for the source-coverage record."""
    connection = connect_current(sqlite_path)
    if connection is None:
        return {}
    doc_types = LARGE_HOLDING_DOC_TYPES + TENDER_OFFER_DOC_TYPES
    try:
        rows = connection.execute(
            "SELECT doc_type_code, COUNT(*) FROM edinet_documents "
            f"WHERE doc_date <= ? AND doc_type_code IN ({_placeholders(doc_types)}) "  # nosec B608
            "GROUP BY doc_type_code ORDER BY doc_type_code",
            (through.isoformat(), *doc_types),
        ).fetchall()
    finally:
        connection.close()
    return {str(code): int(count) for code, count in rows}


def _policy_by_ticker(
    connection: sqlite3.Connection, policy_month: object, requested: Sequence[str]
) -> dict[str, tuple[CapitalPolicyStatus, date | None]]:
    if policy_month is None:
        return {}
    rows = connection.execute(
        "SELECT ticker, status, updated_on FROM tse_capital_policy_snapshots "
        f"WHERE snapshot_month_end = ? AND ticker IN ({_placeholders(requested)})",  # nosec B608
        (str(policy_month), *requested),
    ).fetchall()
    return {
        str(ticker): (_capital_policy_status_from_store(status), _optional_iso_date(updated_on))
        for ticker, status, updated_on in rows
    }


class _TableLinkParser(HTMLParser):
    """Collect table rows and navigation targets from one JPX delisting page."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.rows: list[tuple[str, ...]] = []
        self.links: list[str] = []
        self._cells: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attrs_by_name = dict(attrs)
        if tag == "a" and attrs_by_name.get("href"):
            self.links.append(str(attrs_by_name["href"]))
        if tag == "option" and attrs_by_name.get("value"):
            self.links.append(str(attrs_by_name["value"]))
        if tag == "tr":
            self._cells = []
        elif tag in {"th", "td"} and self._cells is not None:
            self._cell = []

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag in {"th", "td"} and self._cell is not None and self._cells is not None:
            self._cells.append(_text("".join(self._cell)))
            self._cell = None
        elif tag == "tr" and self._cells is not None:
            if self._cells:
                self.rows.append(tuple(self._cells))
            self._cells = None


def _download_jpx_delistings(session: requests.Session) -> tuple[JPXDelistingRow, ...]:
    index_response = session.get(JPX_DELISTING_INDEX_URL, timeout=_HTTP_TIMEOUT_SECONDS)
    if index_response.status_code >= 400:
        raise CapitalControlError(
            f"failed to download JPX delistings: {index_response.status_code}"
        )
    index_html = _decode_jpx_html(index_response)
    index_parser = _TableLinkParser()
    index_parser.feed(index_html)
    urls = {JPX_DELISTING_INDEX_URL}
    for href in index_parser.links:
        resolved = urljoin(JPX_DELISTING_INDEX_URL, href)
        if resolved.startswith(JPX_DELISTING_INDEX_URL) and _JPX_ARCHIVE_RE.search(resolved):
            urls.add(resolved)
    rows: dict[tuple[date, str], JPXDelistingRow] = {}
    for url in sorted(urls):
        response = (
            index_response
            if url == JPX_DELISTING_INDEX_URL
            else session.get(url, timeout=_HTTP_TIMEOUT_SECONDS)
        )
        if response.status_code >= 400:
            raise CapitalControlError(f"failed to download JPX delistings: {url}")
        parser = _TableLinkParser()
        parser.feed(_decode_jpx_html(response))
        for cells in parser.rows:
            row = _jpx_delisting_row(cells)
            if row is None:
                continue
            key = (row.delisted_on, row.ticker)
            previous = rows.get(key)
            if previous is not None and previous != row:
                raise CapitalControlError(
                    f"conflicting JPX delisting row: {row.ticker} {row.delisted_on.isoformat()}"
                )
            rows[key] = row
    if not rows:
        raise CapitalControlError("JPX delisting history had no typed rows")
    return tuple(rows[key] for key in sorted(rows))


def _jpx_delisting_row(cells: Sequence[str]) -> JPXDelistingRow | None:
    """Read one table row, skipping headers and any row that is not a delisting."""
    if len(cells) != 5:
        return None
    try:
        delisted_on = date.fromisoformat(cells[0].replace("/", "-"))
        ticker = normalize_ticker(cells[2])
    except ValueError:
        return None
    if not cells[1] or not cells[4]:
        return None
    return JPXDelistingRow(
        delisted_on=delisted_on,
        ticker=ticker,
        name=cells[1],
        market=cells[3] or None,
        reason=cells[4],
    )


def _decode_jpx_html(response: requests.Response) -> str:
    """Decode by trying strict codecs, never a single-byte one that cannot fail.

    The archive pages arrive without a charset header, and `response.encoding` then
    falls back to latin-1, which decodes every byte sequence into mojibake instead of
    raising. Only codecs that reject invalid input can identify the real encoding.
    """
    for encoding in ("utf-8-sig", "utf-8", "cp932"):
        try:
            return response.content.decode(encoding)
        except (LookupError, UnicodeDecodeError):
            continue
    raise CapitalControlError("failed to decode JPX HTML")


@dataclass(frozen=True, slots=True, kw_only=True)
class _TSEColumns:
    """Where each field sits on one monthly sheet, resolved from its own header."""

    first_data_row: int
    ticker: int
    status: int
    status_change: int
    updated_on: int
    contact_requested: int


# The header spans two rows: a group label with merged cells above, and the label of
# each column inside the group below. Both are needed because a group label alone is
# ambiguous — the contact block prints three columns under one heading, and the column
# holding the disclosure text was inserted before that block in later sheets, so a
# fixed index reads free-form disclosure text as a contact request.
_TSE_TICKER_HEADER = "証券コード"
_TSE_STATUS_GROUP = "開示状況"
_TSE_STATUS_HEADER = "要請に基づく開示状況"
_TSE_STATUS_CHANGE_HEADER = "前月からの開示状況の変更"
_TSE_UPDATED_ON_HEADER = "開示内容のアップデート日"
_TSE_CONTACT_GROUP_PREFIX = "機関投資家からのより活発なコンタクトを希望"
_TSE_CONTACT_HEADER = "申請状況"


def _resolve_tse_columns(rows: Sequence[Sequence[object]], *, sheet: str) -> _TSEColumns:
    header_index = next(
        (
            index
            for index, values in enumerate(rows[:20])
            for labels in ({_text(value) for value in values if value is not None},)
            if _TSE_TICKER_HEADER in labels and _TSE_STATUS_GROUP in labels
        ),
        None,
    )
    if header_index is None or header_index + 1 >= len(rows):
        raise CapitalControlError(f"TSE capital-policy sheet has no typed header: {sheet}")
    group = _forward_filled_labels(rows[header_index])
    detail = [_text(value) for value in rows[header_index + 1]]
    width = max(len(group), len(detail))

    def find(predicate: Callable[[int], bool], label: str) -> int:
        column = next((index for index in range(width) if predicate(index)), None)
        if column is None:
            raise CapitalControlError(f"TSE capital-policy sheet has no {label} column: {sheet}")
        return column

    def at(labels: Sequence[str], index: int) -> str:
        return labels[index] if index < len(labels) else ""

    return _TSEColumns(
        first_data_row=header_index + 2,
        ticker=find(lambda index: at(group, index) == _TSE_TICKER_HEADER, _TSE_TICKER_HEADER),
        status=find(lambda index: at(detail, index) == _TSE_STATUS_HEADER, _TSE_STATUS_HEADER),
        status_change=find(
            lambda index: at(detail, index) == _TSE_STATUS_CHANGE_HEADER,
            _TSE_STATUS_CHANGE_HEADER,
        ),
        updated_on=find(
            lambda index: at(group, index) == _TSE_UPDATED_ON_HEADER, _TSE_UPDATED_ON_HEADER
        ),
        contact_requested=find(
            lambda index: (
                at(group, index).startswith(_TSE_CONTACT_GROUP_PREFIX)
                and at(detail, index) == _TSE_CONTACT_HEADER
            ),
            _TSE_CONTACT_GROUP_PREFIX,
        ),
    )


def _forward_filled_labels(values: Sequence[object]) -> list[str]:
    """Spread a merged group label across the columns it covers."""
    labels: list[str] = []
    current = ""
    for value in values:
        text = _text(value)
        if text:
            current = text
        labels.append(current)
    return labels


def _capital_policy_status(value: str) -> CapitalPolicyStatus:
    if value == "開示済":
        return "disclosed"
    if value == "検討中":
        return "considering"
    raise CapitalControlError(f"unknown TSE capital-policy status: {value!r}")


def _capital_policy_status_from_store(value: object) -> CapitalPolicyStatus:
    raw = str(value)
    if raw == "disclosed":
        return "disclosed"
    if raw == "considering":
        return "considering"
    raise CapitalControlError(f"invalid stored TSE capital-policy status: {raw!r}")


def _ticker_from_excel(value: object) -> str:
    if isinstance(value, bool):
        raise CapitalControlError(f"invalid TSE ticker: {value!r}")
    if isinstance(value, int):
        raw = str(value)
    elif isinstance(value, float) and value.is_integer():
        raw = str(int(value))
    else:
        raw = _text(value)
    try:
        return normalize_ticker(raw)
    except ValueError as exc:
        raise CapitalControlError(f"invalid TSE ticker: {value!r}") from exc


def _excel_date(value: object) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = _text(value)
    try:
        return date.fromisoformat(
            raw.replace("/", "-").replace("年", "-").replace("月", "-").removesuffix("日")
        )
    except ValueError as exc:
        raise CapitalControlError(f"invalid TSE update date: {raw!r}") from exc


def _calendar_month_end(year: int, month: int) -> date:
    first_next = date(year + int(month == 12), month % 12 + 1, 1)
    return first_next - timedelta(days=1)


def _cell(values: Sequence[object], index: int) -> object | None:
    return values[index] if index < len(values) else None


def _text(value: object) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip()


def _status_text(value: object) -> str | None:
    return None if value is None else str(value)


def _placeholders(values: Sequence[object]) -> str:
    return ",".join("?" for _ in values)


def _date_iso(value: date | None) -> str | None:
    return None if value is None else value.isoformat()


def _optional_iso_date(value: object) -> date | None:
    return None if value in (None, "") else date.fromisoformat(str(value))


__all__ = (
    "EVENT_RECENT_DAYS",
    "LARGE_HOLDING_DOC_TYPES",
    "TENDER_OFFER_DOC_TYPES",
    "TENDER_OFFER_REGISTRATION_CORRECTION_DOC_TYPE",
    "TENDER_OFFER_REGISTRATION_DOC_TYPE",
    "TENDER_OFFER_RESULT_CORRECTION_DOC_TYPE",
    "TENDER_OFFER_RESULT_DOC_TYPE",
    "TENDER_OFFER_WITHDRAWAL_DOC_TYPE",
    "CapitalControlAnnotation",
    "CapitalControlError",
    "CapitalControlRefreshSummary",
    "CapitalPolicyStatus",
    "ControlEventIndex",
    "ControlEventType",
    "JPXDelistingRow",
    "TSECapitalPolicyRow",
    "edinet_code_to_ticker",
    "edinet_event_counts",
    "edinet_identity_covered",
    "is_usable_filing_status",
    "parse_tse_capital_policy_workbook",
    "read_capital_control_annotations",
    "read_control_event_index",
    "read_jpx_delistings",
    "refresh_capital_control_facts",
    "store_jpx_delistings",
    "store_tse_capital_policy_rows",
)
