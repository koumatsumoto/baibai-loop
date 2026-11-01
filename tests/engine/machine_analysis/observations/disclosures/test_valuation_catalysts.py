from __future__ import annotations

import unittest
from contextlib import closing
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

import openpyxl
import pytest

import baibai_engine.machine_analysis.observations.disclosures.valuation_catalysts as catalysts
from baibai_engine.machine_analysis.observations.disclosures.valuation_catalysts import (
    ValuationCatalystError,
    edinet_code_to_ticker,
    parse_tse_capital_policy_workbook,
)
from baibai_engine.storage.sqlite.source_cache.schema import open_connection


# The two real layouts. The 2025 sheets put the contact application next to the update
# date; the 2026 sheets inserted the disclosure text between them, which is what makes a
# fixed column index read free-form prose as a contact request.
def _naive_datetime(year: int, month: int, day: int) -> datetime:
    """openpyxl writes cell dates without a zone, which is what the reader receives."""
    return datetime(year, month, day)  # noqa: DTZ001


_GROUP_2025 = (
    None,
    "業種\nコード",
    "業種",
    "市場区分",
    "証券\nコード",
    "銘柄名",
    "開示状況",
    None,
    "開示内容の\nアップデート日",
    "機関投資家からのより活発なコンタクトを希望",
    None,
    None,
    "英文開示",
)
_GROUP_2026 = (
    None,
    "業種\nコード",
    "業種",
    "市場区分",
    "証券\nコード",
    "銘柄名",
    "開示状況",
    None,
    "開示内容の\nアップデート日",
    "開示内容 ※１",
    "機関投資家からのより活発なコンタクトを希望 ※２",
    None,
    None,
    "英文開示",
)
_DETAIL_2025 = (
    None,
    None,
    None,
    None,
    None,
    None,
    "要請に基づく\n開示状況",
    "前月からの\n開示状況の変更",
    None,
    "申請状況",
    "掲載開始日",
    "コンタクト先",
    None,
)
_DETAIL_2026 = (
    None,
    None,
    None,
    None,
    None,
    None,
    "要請に基づく\n開示状況",
    "前月からの\n開示状況の変更",
    None,
    None,
    "申請状況",
    "掲載開始日",
    "コンタクト先",
    None,
)


def _workbook(
    sheets: list[tuple[str, tuple[object, ...], tuple[object, ...], list[tuple]]],
) -> bytes:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for title, group, detail, rows in sheets:
        sheet = workbook.create_sheet(title=title)
        sheet.append(("本資料の目的",))
        sheet.append(group)
        sheet.append(detail)
        for row in rows:
            sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _row_2026(ticker: int, status: str, *, updated: object = "", contact: str = "") -> tuple:
    return (
        None,
        50,
        "水産・農林業",
        "プライム",
        ticker,
        "極洋",
        status,
        "",
        updated,
        "本文",
        contact,
        "",
        "",
        "有",
    )


def _row_2025(ticker: int, status: str, *, updated: object = "", contact: str = "") -> tuple:
    return (
        None,
        50,
        "水産・農林業",
        "プライム",
        ticker,
        "極洋",
        status,
        "",
        updated,
        contact,
        "",
        "",
        "有",
    )


class TSECapitalPolicyWorkbookTest(unittest.TestCase):
    def test_columns_are_resolved_per_sheet_rather_than_by_position(self) -> None:
        content = _workbook(
            [
                (
                    "開示企業一覧（2026年6月末時点）",
                    _GROUP_2026,
                    _DETAIL_2026,
                    [_row_2026(1301, "開示済", updated=_naive_datetime(2026, 6, 23), contact="○")],
                ),
                (
                    "【過去分】開示企業一覧（2025年5月末時点）",
                    _GROUP_2025,
                    _DETAIL_2025,
                    [_row_2025(1301, "検討中")],
                ),
            ]
        )

        rows = {row.snapshot_month_end: row for row in parse_tse_capital_policy_workbook(content)}

        latest = rows[date(2026, 6, 30)]
        self.assertEqual(latest.status, "disclosed")
        self.assertEqual(latest.updated_on, date(2026, 6, 23))
        self.assertTrue(latest.contact_requested)
        earliest = rows[date(2025, 5, 31)]
        self.assertEqual(earliest.status, "considering")
        self.assertFalse(earliest.contact_requested)

    def test_disclosure_text_is_not_read_as_a_contact_request(self) -> None:
        content = _workbook(
            [
                (
                    "開示企業一覧（2026年6月末時点）",
                    _GROUP_2026,
                    _DETAIL_2026,
                    [_row_2026(1301, "開示済")],
                )
            ]
        )

        (row,) = parse_tse_capital_policy_workbook(content)

        self.assertFalse(row.contact_requested)

    def test_first_disclosure_on_the_earliest_sheet_is_left_censored(self) -> None:
        content = _workbook(
            [
                (
                    "開示企業一覧（2026年6月末時点）",
                    _GROUP_2026,
                    _DETAIL_2026,
                    [_row_2026(1301, "開示済"), _row_2026(1332, "開示済")],
                ),
                (
                    "【過去分】開示企業一覧（2025年5月末時点）",
                    _GROUP_2025,
                    _DETAIL_2025,
                    [_row_2025(1301, "開示済"), _row_2025(1332, "検討中")],
                ),
            ]
        )

        rows = {
            (row.ticker, row.snapshot_month_end): row
            for row in parse_tse_capital_policy_workbook(content)
        }

        already_listed = rows[("1301", date(2026, 6, 30))]
        self.assertEqual(already_listed.first_disclosed_month_end, date(2025, 5, 31))
        self.assertTrue(already_listed.first_disclosure_left_censored)
        newly_disclosed = rows[("1332", date(2026, 6, 30))]
        self.assertEqual(newly_disclosed.first_disclosed_month_end, date(2026, 6, 30))
        self.assertFalse(newly_disclosed.first_disclosure_left_censored)
        earlier = rows[("1332", date(2025, 5, 31))]
        self.assertIsNone(earlier.first_disclosed_month_end)
        self.assertFalse(earlier.first_disclosure_left_censored)

    def test_invalid_month_is_not_normalized_into_another_month(self) -> None:
        for month in (0, 13):
            with self.subTest(month=month):
                content = _workbook(
                    [
                        (
                            f"開示企業一覧（2026年{month}月末時点）",
                            _GROUP_2026,
                            _DETAIL_2026,
                            [_row_2026(1301, "開示済")],
                        )
                    ]
                )
                with self.assertRaisesRegex(ValuationCatalystError, "invalid TSE snapshot month"):
                    parse_tse_capital_policy_workbook(content)

    def test_month_end_handles_leap_year_and_december(self) -> None:
        for expected in (date(2024, 2, 29), date(2026, 12, 31)):
            with self.subTest(expected=expected):
                content = _workbook(
                    [
                        (
                            f"開示企業一覧（{expected.year}年{expected.month}月末時点）",
                            _GROUP_2026,
                            _DETAIL_2026,
                            [_row_2026(1301, "開示済")],
                        )
                    ]
                )
                (row,) = parse_tse_capital_policy_workbook(content)
                self.assertEqual(row.snapshot_month_end, expected)

    def test_an_unknown_status_word_is_refused_rather_than_guessed(self) -> None:
        content = _workbook(
            [
                (
                    "開示企業一覧（2026年6月末時点）",
                    _GROUP_2026,
                    _DETAIL_2026,
                    [_row_2026(1301, "未定")],
                )
            ]
        )

        with self.assertRaises(ValuationCatalystError):
            parse_tse_capital_policy_workbook(content)

    def test_a_sheet_without_the_typed_header_is_refused(self) -> None:
        content = _workbook(
            [
                (
                    "開示企業一覧（2026年6月末時点）",
                    (None, "銘柄", "市場"),
                    (None, None, None),
                    [(None, 1301, "プライム")],
                )
            ]
        )

        with self.assertRaises(ValuationCatalystError):
            parse_tse_capital_policy_workbook(content)

    def test_a_workbook_with_no_monthly_sheet_is_refused(self) -> None:
        content = _workbook([("本資料の目的", (None,), (None,), [])])

        with self.assertRaises(ValuationCatalystError):
            parse_tse_capital_policy_workbook(content)


def _store_documents(sqlite_path: Path, rows: list[tuple[object, ...]]) -> None:
    connection = open_connection(sqlite_path)
    connection.executemany(
        "INSERT INTO edinet_documents("
        "doc_date, sequence_number, doc_id, sec_code, doc_type_code, legal_status, "
        "disclosure_status, withdrawal_status, edinet_code, issuer_edinet_code, "
        "subject_edinet_code) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        rows,
    )
    connection.executemany(
        "INSERT OR REPLACE INTO edinet_document_lists("
        "doc_date, process_datetime, result_count, selected_received_at_utc, "
        "latest_received_at_utc, rechecked_asof"
        ") VALUES (?, NULL, 1, '2026-08-11T00:00:00+00:00', "
        "'2026-08-11T00:00:00+00:00', '2026-08-11')",
        sorted({(str(row[0]),) for row in rows}),
    )
    connection.commit()
    connection.close()


class EdinetCodeMappingTest(unittest.TestCase):
    ASOF = date(2026, 8, 10)

    def test_identity_mapping_does_not_truncate_nonordinary_or_malformed_codes(self) -> None:
        with TemporaryDirectory() as tmp:
            sqlite_path = Path(tmp) / "market.sqlite"
            _store_documents(
                sqlite_path,
                [
                    ("2026-06-01", 1, "S0001", "13010", "120", "1", "0", "0", "E1", None, None),
                    ("2026-06-01", 2, "S0002", "13011", "120", "1", "0", "0", "E2", None, None),
                    ("2026-06-01", 3, "S0003", "1301bad", "120", "1", "0", "0", "E3", None, None),
                    ("2026-06-01", 4, "S0004", "130a0", "120", "1", "0", "0", "E4", None, None),
                ],
            )
            with closing(open_connection(sqlite_path)) as connection:
                self.assertEqual(
                    edinet_code_to_ticker(connection, through=self.ASOF),
                    {"E1": "1301", "E4": "130A"},
                )

    def test_future_identity_and_conflict_do_not_change_historical_mapping(self) -> None:
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "market.sqlite"
            _store_documents(
                path,
                [
                    ("2026-06-01", 1, "S0001", "13010", "120", "1", "0", "0", "E1", None, None),
                    ("2026-08-03", 1, "S0002", None, "350", "1", "0", "0", "E9", "E1", None),
                    ("2026-08-11", 1, "S0003", "13020", "120", "1", "0", "0", "E1", None, None),
                    ("2026-08-11", 2, "S0004", "13030", "120", "1", "0", "0", "E3", None, None),
                ],
            )
            with closing(open_connection(path)) as connection:
                before = edinet_code_to_ticker(connection, through=self.ASOF)
                after = edinet_code_to_ticker(connection, through=date(2026, 8, 11))
            self.assertEqual(before, {"E1": "1301"})
            self.assertEqual(after, {"E3": "1303"})


def test_workbook_is_closed_when_monthly_rows_are_missing(mocker):
    workbook = mocker.Mock(worksheets=[])
    mocker.patch.object(catalysts.openpyxl, "load_workbook", return_value=workbook)
    with pytest.raises(catalysts.ValuationCatalystError, match="no monthly rows"):
        catalysts.parse_tse_capital_policy_workbook(b"workbook")
    workbook.close.assert_called_once_with()


if __name__ == "__main__":
    unittest.main()
