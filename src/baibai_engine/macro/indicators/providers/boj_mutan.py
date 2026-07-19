from __future__ import annotations

import io
import re
import zipfile
from datetime import date
from html import unescape
from urllib.parse import urljoin

import openpyxl

from ..db import ObservationRecord
from ..definitions import SeriesDefinition
from .base import (
    MAX_CSV_RESPONSE_BYTES,
    MAX_ZIP_RESPONSE_BYTES,
    FetchContext,
    HttpSession,
    IndicatorsProviderError,
    fetch_bytes,
    fetch_text,
    parse_float,
    record_observation,
)

_CURRENT_XLSX_START = date(2025, 10, 1)
_OLD_ARCHIVE_END = date(2025, 9, 30)
_OLD_MENU_URL_TEMPLATE = "https://www3.boj.or.jp/market/jp/menuold_m_{year}.htm"
_FINAL_XLSX_RE = re.compile(r'href="(?P<href>[^"]*md(?P<date>\d{8})\.xlsx)"')
_OLD_HTML_RE = re.compile(
    r"""href=["'](?P<href>[^"']*md(?P<date>\d{6})\.htm)["']""",
    re.IGNORECASE,
)
_OLD_AVERAGE_RE = re.compile(r"平均\s*(?P<value>[+-]?\d+(?:,\d{3})*(?:\.\d+)?)\s*[％%]")


class BojMutanProvider:
    """BOJ daily final uncollateralized overnight call rate.

    BOJ publishes one XLSX per business day. The final result file name carries
    the observation date as ``mdYYYYMMDD.xlsx``; the provider reads the weighted
    average row from each final workbook.
    """

    name = "boj_mutan"

    def fetch(
        self,
        series: SeriesDefinition,
        *,
        start: date,
        end: date,
        session: HttpSession,
        context: FetchContext | None = None,
    ) -> list[ObservationRecord]:
        observations: list[ObservationRecord] = []
        if end >= _CURRENT_XLSX_START:
            current_start = max(start, _CURRENT_XLSX_START)
            for observed_at, url in _final_xlsx_links(
                series.source_url,
                start=current_start,
                end=end,
                session=session,
                context=context,
            ):
                content = fetch_bytes(
                    session,
                    url,
                    params=None,
                    max_bytes=MAX_ZIP_RESPONSE_BYTES,
                    context=context,
                )
                observations.append(parse_boj_mutan_xlsx(series, content, observed_at=observed_at))
        if start <= _OLD_ARCHIVE_END:
            old_end = min(end, _OLD_ARCHIVE_END)
            for observed_at, url in _old_html_links(
                start=start,
                end=old_end,
                session=session,
                context=context,
            ):
                content = fetch_bytes(
                    session,
                    url,
                    params=None,
                    max_bytes=MAX_CSV_RESPONSE_BYTES,
                    context=context,
                )
                observations.append(
                    parse_boj_mutan_old_html(series, content, observed_at=observed_at)
                )
        return sorted(observations, key=lambda item: item.observed_at)


def parse_boj_mutan_xlsx(
    series: SeriesDefinition, content: bytes, *, observed_at: date
) -> ObservationRecord:
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise IndicatorsProviderError("BOJ MUTAN response is not a .xlsx (zip) workbook")
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            worksheet = workbook[workbook.sheetnames[0]]
            for row in worksheet.iter_rows(values_only=True):
                values = list(row)
                for index, cell in enumerate(values):
                    if isinstance(cell, str) and "Average" in cell:
                        value = _next_numeric(values[index + 1 :])
                        if value is None:
                            raise IndicatorsProviderError("BOJ MUTAN Average row has no value")
                        return record_observation(series, observed_at=observed_at, value=value)
        finally:
            workbook.close()
    except IndicatorsProviderError:
        raise
    except Exception as exc:
        raise IndicatorsProviderError(f"BOJ MUTAN workbook could not be read: {exc}") from exc
    raise IndicatorsProviderError("BOJ MUTAN workbook missing Average row")


def parse_boj_mutan_old_html(
    series: SeriesDefinition, content: bytes, *, observed_at: date
) -> ObservationRecord:
    return record_observation(
        series,
        observed_at=observed_at,
        value=parse_boj_mutan_old_average(content),
    )


def parse_boj_mutan_old_average(content: bytes) -> float:
    for text in _decode_html_candidates(content):
        normalized = unescape(re.sub(r"<[^>]+>", " ", text))
        normalized = re.sub(r"\s+", " ", normalized)
        match = _OLD_AVERAGE_RE.search(normalized)
        if match:
            return parse_float(match.group("value"))
    raise IndicatorsProviderError("BOJ MUTAN old archive page missing average")


def _final_xlsx_links(
    source_url: str,
    *,
    start: date,
    end: date,
    session: HttpSession,
    context: FetchContext | None,
) -> list[tuple[date, str]]:
    links: dict[date, str] = {}
    for year in range(start.year, end.year + 1):
        index_url = (
            f"https://www.boj.or.jp/statistics/market/short/mutan/d_release/md/{year}/index.htm"
        )
        html = fetch_text(
            session,
            index_url,
            params=None,
            max_bytes=MAX_CSV_RESPONSE_BYTES,
            context=context,
        )
        for match in _FINAL_XLSX_RE.finditer(html):
            observed_at = date.fromisoformat(
                f"{match.group('date')[:4]}-{match.group('date')[4:6]}-{match.group('date')[6:]}"
            )
            if start <= observed_at <= end:
                links[observed_at] = urljoin(index_url, match.group("href"))
    if not links and start.year == end.year:
        html = fetch_text(
            session,
            source_url,
            params=None,
            max_bytes=MAX_CSV_RESPONSE_BYTES,
            context=context,
        )
        for match in _FINAL_XLSX_RE.finditer(html):
            observed_at = date.fromisoformat(
                f"{match.group('date')[:4]}-{match.group('date')[4:6]}-{match.group('date')[6:]}"
            )
            if start <= observed_at <= end:
                links[observed_at] = urljoin(source_url, match.group("href"))
    return sorted(links.items())


def _old_html_links(
    *,
    start: date,
    end: date,
    session: HttpSession,
    context: FetchContext | None,
) -> list[tuple[date, str]]:
    links: dict[date, str] = {}
    for year in range(start.year, end.year + 1):
        menu_url = _OLD_MENU_URL_TEMPLATE.format(year=year)
        html = _decode_html_for_links(
            fetch_bytes(
                session,
                menu_url,
                params=None,
                max_bytes=MAX_CSV_RESPONSE_BYTES,
                context=context,
            )
        )
        for match in _OLD_HTML_RE.finditer(html):
            raw_date = match.group("date")
            observed_at = date.fromisoformat(f"20{raw_date[:2]}-{raw_date[2:4]}-{raw_date[4:]}")
            if start <= observed_at <= end:
                links[observed_at] = urljoin(menu_url, match.group("href"))
    return sorted(links.items())


def _decode_html_for_links(content: bytes) -> str:
    return next(
        iter(_decode_html_candidates(content)),
        content.decode("utf-8-sig", errors="replace"),
    )


def _decode_html_candidates(content: bytes) -> tuple[str, ...]:
    candidates: list[str] = []
    for encoding in ("cp932", "utf-8-sig", "utf-8"):
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError:
            continue
        if text not in candidates:
            candidates.append(text)
    return tuple(candidates)


def _next_numeric(cells: list[object]) -> float | None:
    for cell in cells:
        if isinstance(cell, bool):
            continue
        if isinstance(cell, (int, float)):
            return float(cell)
    return None
