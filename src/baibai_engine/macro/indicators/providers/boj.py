from __future__ import annotations

import io
import math
import zipfile
from collections.abc import Sequence
from datetime import date, datetime

import openpyxl

from ..db import ObservationRecord
from ..definitions import SeriesDefinition
from .base import (
    MAX_ZIP_RESPONSE_BYTES,
    FetchContext,
    HttpSession,
    IndicatorsProviderError,
    ProviderSpec,
    fetch_bytes,
    record_observation,
)


class BojProvider:
    """Bank of Japan long time-series Excel workbook (no auth).

    BOJ publishes stable ``.xlsx`` long-series workbooks (e.g. the monetary base
    at ``other/mb/mblong.xlsx``); the stat-search interactive CSV is session /
    cgi based and not a stable GET URL. The first sheet holds the headline series
    with a month-end date in the date column and values in numbered data columns.
    ``provider_series_id`` binds the 1-based value column and header to a
    positioned workbook scale marker (e.g.
    ``"3|Monetary Base|8|Unit: 100 million yen"``). All are checked so a workbook
    column insertion or base/unit change cannot silently alter the observations.
    Observations are normalised to the first of the month so they align with the
    other monthly series (FRED / e-Stat).
    """

    spec = ProviderSpec(name="boj", all_history_start=date(1957, 1, 1))
    name = spec.name

    def fetch(
        self,
        series: SeriesDefinition,
        *,
        start: date,
        end: date,
        session: HttpSession,
        context: FetchContext | None = None,
    ) -> list[ObservationRecord]:
        content = fetch_bytes(
            session,
            series.source_url,
            params=None,
            max_bytes=MAX_ZIP_RESPONSE_BYTES,
            context=context,
        )
        return parse_boj_xlsx(series, content, start=start, end=end)


def parse_boj_xlsx(
    series: SeriesDefinition, content: bytes, *, start: date, end: date
) -> list[ObservationRecord]:
    value_col, expected_header, metadata_col, expected_metadata = _value_column_spec(
        series.provider_series_id
    )
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise IndicatorsProviderError("BOJ response is not a .xlsx (zip) workbook")
    observations: list[ObservationRecord] = []
    saw_date_row = False
    in_range_date_rows = 0
    header_cells: list[str] = []
    metadata_cells: list[str] = []
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            worksheet = workbook[workbook.sheetnames[0]]
            for row in worksheet.iter_rows(values_only=True):
                cells: Sequence[object] = row
                observed_at = _row_month(cells)
                if observed_at is None:
                    if not saw_date_row:
                        if metadata_col - 1 < len(cells):
                            metadata = cells[metadata_col - 1]
                            if isinstance(metadata, str) and metadata.strip():
                                metadata_cells.append(_normalize_header(metadata))
                        if value_col - 1 < len(cells):
                            header = cells[value_col - 1]
                            if isinstance(header, str) and header.strip():
                                header_cells.append(_normalize_header(header))
                    continue
                saw_date_row = True
                if not start <= observed_at <= end:
                    continue
                in_range_date_rows += 1
                cell = cells[value_col - 1] if value_col - 1 < len(cells) else None
                value = _cell_float(cell)
                if value is None:
                    continue
                observations.append(
                    record_observation(series, observed_at=observed_at, value=value)
                )
        finally:
            workbook.close()
    except Exception as exc:
        # openpyxl / ElementTree raise opaque third-party errors on a corrupt or
        # non-xlsx workbook (InvalidFileException, ParseError, IndexError on an
        # empty workbook); convert them so the indicators CLI never leaks a traceback.
        raise IndicatorsProviderError(f"BOJ workbook could not be read: {exc}") from exc
    if not saw_date_row:
        raise IndicatorsProviderError("BOJ workbook has no parseable date rows")
    normalized_expected = _normalize_header(expected_header)
    if normalized_expected not in header_cells:
        raise IndicatorsProviderError(
            f"BOJ value column {value_col} header mismatch: expected {expected_header!r}"
        )
    normalized_metadata = _normalize_header(expected_metadata)
    if normalized_metadata not in metadata_cells:
        raise IndicatorsProviderError(
            f"BOJ metadata column {metadata_col} mismatch: expected {expected_metadata!r}"
        )
    if in_range_date_rows and not observations:
        raise IndicatorsProviderError(
            f"BOJ value column {value_col} has {in_range_date_rows} date rows "
            "in the requested range but no numeric values"
        )
    return observations


def _value_column_spec(provider_series_id: str) -> tuple[int, str, int, str]:
    parts = provider_series_id.split("|")
    if len(parts) != 4 or not parts[1].strip() or not parts[3].strip():
        raise IndicatorsProviderError(
            "BOJ provider_series_id must be "
            "'<value column>|<expected header>|<metadata column>|<expected metadata>': "
            f"{provider_series_id!r}"
        )
    raw_index, expected_header, raw_metadata_index, expected_metadata = parts
    try:
        index = int(raw_index)
    except ValueError:
        raise IndicatorsProviderError(
            f"BOJ provider_series_id must start with a 1-based column index: {provider_series_id!r}"
        ) from None
    if index < 2:
        raise IndicatorsProviderError(
            f"BOJ value column index must be a 1-based value column (>= 2): {index}"
        )
    try:
        metadata_index = int(raw_metadata_index)
    except ValueError:
        raise IndicatorsProviderError(
            "BOJ provider_series_id metadata column must be a 1-based column index: "
            f"{provider_series_id!r}"
        ) from None
    if metadata_index < 1:
        raise IndicatorsProviderError(
            f"BOJ metadata column index must be 1-based (>= 1): {metadata_index}"
        )
    return (
        index,
        expected_header.strip(),
        metadata_index,
        expected_metadata.strip(),
    )


def _normalize_header(value: str) -> str:
    return " ".join(value.split()).casefold()


def _row_month(cells: Sequence[object]) -> date | None:
    for cell in cells:
        if isinstance(cell, datetime):
            return date(cell.year, cell.month, 1)
        if isinstance(cell, date):
            return date(cell.year, cell.month, 1)
    return None


def _cell_float(cell: object) -> float | None:
    if isinstance(cell, bool):
        return None
    if isinstance(cell, (int, float)):
        value = float(cell)
        return value if math.isfinite(value) else None
    return None
